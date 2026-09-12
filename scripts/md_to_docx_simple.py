import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")


def set_base_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)


def add_text_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(11 if style != "Title" else 18)


def add_image(document: Document, image_path: Path) -> None:
    if not image_path.exists():
        add_text_paragraph(document, f"[图片缺失] {image_path}")
        return
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.2))


def handle_line(document: Document, raw_line: str) -> None:
    line = raw_line.rstrip()
    stripped = line.strip()
    if not stripped:
        return

    image_match = IMAGE_RE.search(stripped)
    if image_match:
        image_path = Path(image_match.group(1))
        add_image(document, image_path)
        return

    if stripped.startswith("# "):
        add_text_paragraph(document, stripped[2:].strip(), "Title")
        return

    if stripped.startswith("## "):
        add_text_paragraph(document, stripped[3:].strip(), "Heading 1")
        return

    if stripped.startswith("### "):
        add_text_paragraph(document, stripped[4:].strip(), "Heading 2")
        return

    bold_heading = re.fullmatch(r"\*\*(.+?)\*\*", stripped)
    if bold_heading:
        add_text_paragraph(document, bold_heading.group(1).strip(), "Heading 1")
        return

    if stripped.startswith("- "):
        add_text_paragraph(document, stripped[2:].strip(), "List Bullet")
        return

    add_text_paragraph(document, stripped)


def convert(md_path: Path, output_path: Path) -> None:
    content = md_path.read_text(encoding="utf-8")
    document = Document()
    set_base_style(document)

    for line in content.splitlines():
        handle_line(document, line)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 3:
        raise SystemExit("Usage: md_to_docx_simple.py <input.md> <output.docx>")

    convert(Path(sys.argv[1]), Path(sys.argv[2]))
