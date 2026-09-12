from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt


def set_run_font(run, name: str, size_pt: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size_pt)
    r_fonts = run._element.rPr.rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def insert_paragraph_before(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_heading_para(paragraph: Paragraph, text: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, "黑体", 12, bold=True)


def add_body_para(paragraph: Paragraph, text: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 12)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    doc = Document(args.input)

    target = None
    for para in doc.paragraphs:
        if para.text.strip() == "6.6 程序改进设想":
            target = para
            break

    if target is None:
        raise RuntimeError("Could not find section 6.6 heading")

    p1 = insert_paragraph_before(target)
    add_heading_para(p1, "6.6 已报名竞赛页面性能调试")

    p2 = insert_paragraph_before(target)
    add_body_para(
        p2,
        "问题描述：学生端“已报名竞赛”页面首次打开时存在明显延迟，抓包后发现页面组件与导航栏预取逻辑同时请求报名记录接口，首次进入通常需要约 5s~10s，影响用户体验。",
    )

    p3 = insert_paragraph_before(target)
    add_body_para(
        p3,
        "原因分析：原有查询链路存在全量查询与业务层过滤并存的问题，报名记录、参赛记录和评审状态需要逐条组装；同时，数据库连接池的失活连接、通用组件首次初始化以及接口内部的 N+1 查询都会放大首屏耗时。",
    )

    p4 = insert_paragraph_before(target)
    add_body_para(
        p4,
        "解决方案：将页面改为独立分页接口，减少重复请求；对报名状态采用批量查询与页内预加载；优化连接池校验与保活参数；将通用组件初始化前置到启动阶段。调整后，页面首次加载的耗时得到明显压缩。",
    )

    if target.runs:
        target.runs[0].text = "6.7 程序改进设想"
    else:
        target.text = "6.7 程序改进设想"

    doc.save(args.output)


if __name__ == "__main__":
    main()
